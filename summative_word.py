#!/usr/bin/env python3
"""
summative_word.py
=================
Word (.docx) rendition of the TVET CDACC Summative Assessment Moderated
Practical Marks Sheet — the same template as summative_excel.py, one
document per unit of competency.

Word can't reproduce a spreadsheet pixel-for-pixel, but the structure,
wording, fonts (Times New Roman), the blue title, the grey-shaded table
header and the signatory block all follow the same template.
"""

import io
import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from marksheet_excel import _safe_filename

_HERE       = os.path.dirname(os.path.abspath(__file__))
_CDACC_LOGO = os.path.join(_HERE, "cdacc_logo.png")
_FONT       = "Times New Roman"
_BLUE       = RGBColor(0x1F, 0x4E, 0x79)
_GREY       = "D9D9D9"

_ORG   = "TVET CURRICULUM DEVELOPMENT, ASSESSMENT AND CERTIFICATION COUNCIL (TVET CDACC)"
_TITLE = "SUMMATIVE ASSESSMENT MODERATED PRACTICAL MARKS SHEETS PER UNIT OF COMPETENCY"


def _dots(n: int = 50) -> str:
    return "." * n


def _runs(paragraph, parts, size: int = 11, color: RGBColor = None):
    """Add (bold, text) run pairs to a paragraph in Times New Roman."""
    for bold, text in parts:
        run = paragraph.add_run(text)
        run.bold = bool(bold)
        run.font.name = _FONT
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return paragraph


def _para(doc_or_cell, parts, size=11, align=WD_ALIGN_PARAGRAPH.LEFT,
          color=None, space_after=2):
    p = doc_or_cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(space_after)
    _runs(p, parts, size=size, color=color)
    return p


def _row_height(row, pt: float):
    """Match the template's row heights. AT_LEAST (not EXACT) so wrapped
    content — a long name in Word's narrower columns — is never clipped."""
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    row.height = Pt(pt)


def _shade(cell, fill_hex: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def _no_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tbl_pr.append(borders)


def _cell(cell, parts, size=11, align=WD_ALIGN_PARAGRAPH.LEFT,
          bold_all=False, shade=None, vcenter=True):
    """Fill a table cell: clear default paragraph, add wrapped/​multiline runs."""
    cell.paragraphs[0].text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    for i, (bold, text) in enumerate(parts):
        # honour embedded newlines as line breaks within the cell
        chunks = str(text).split("\n")
        for j, chunk in enumerate(chunks):
            run = p.add_run(chunk)
            run.bold = bool(bold or bold_all)
            run.font.name = _FONT
            run.font.size = Pt(size)
            if j < len(chunks) - 1:
                run.add_break()
    if vcenter:
        cell.vertical_alignment = 1  # WD_ALIGN_VERTICAL.CENTER
    if shade:
        _shade(cell, shade)


def _setup_page(doc):
    s = doc.sections[0]
    s.orientation   = WD_ORIENT.PORTRAIT
    s.page_width    = Mm(210)
    s.page_height   = Mm(297)
    s.left_margin   = Inches(0.5)
    s.right_margin  = Inches(0.5)
    s.top_margin    = Inches(0.5)
    s.bottom_margin = Inches(0.5)
    # default document font
    normal = doc.styles["Normal"]
    normal.font.name = _FONT
    normal.font.size = Pt(11)


# ── Document builder ──────────────────────────────────────────────────────────

def build_summative_doc(doc, data: dict, unit: dict):
    centre_code  = data.get("centre_code") or ""
    centre_name  = data.get("centre_name") or ""
    course_name  = unit.get("course_name") or data.get("course_name") or ""
    course_level = unit.get("course_level") or data.get("course_level") or ""
    unit_code    = unit.get("unit_code") or ""
    unit_name    = unit["unit_name"]
    series       = data.get("series") or ""
    candidates   = sorted(unit["candidates"], key=lambda c: c["sn"])

    lvl = course_level.strip()
    if lvl and not lvl.lower().startswith("level"):
        lvl = f"Level {lvl}"
    course_full = f"{course_name}  {lvl}".strip() if lvl else course_name

    # ── Logo (centred) ────────────────────────────────────────────────────────
    if os.path.exists(_CDACC_LOGO):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(_CDACC_LOGO, height=Inches(0.9))

    # ── Headings ──────────────────────────────────────────────────────────────
    _para(doc, [(True, _ORG)],   size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, [(True, _TITLE)], size=12, align=WD_ALIGN_PARAGRAPH.CENTER, color=_BLUE)
    _para(doc, [(True, "CBA/SAMSP/1")], size=11)

    # ── Info grid (borderless 2-column table) ─────────────────────────────────
    info = doc.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_table_borders(info)
    info_rows = [
        (("Course/Qualification Code:  ", _dots(30)),
         ("Course/Qualification Title:  ", course_full or _dots(30))),
        (("Assessment Center Code:  ", centre_code or _dots(28)),
         ("Assessment Center Name:  ", centre_name or _dots(28))),
        (("Unit Code:  ", unit_code or _dots(30)),
         ("Unit Title:  ", unit_name or _dots(30))),
        (("Date of Assessment:  From ", _dots(14) + "  to  " + _dots(14)),
         ("Assessment Series:  ", series or _dots(26))),
    ]
    for r, (left, right) in enumerate(info_rows):
        _cell(info.rows[r].cells[0], [(True, left[0]), (False, left[1])])
        _cell(info.rows[r].cells[1], [(True, right[0]), (False, right[1])])
        _row_height(info.rows[r], 20)   # template rows 7-10

    # ── Name / mean-deviation lines ───────────────────────────────────────────
    _para(doc, [(True, "1. Name of Internal Assessor:  "), (False, _dots(40)),
                (True, "      Mean Deviation:  "), (False, _dots(16))], size=11)
    _para(doc, [(True, "2. Name of External Verifier:  "), (False, _dots(70))], size=11)
    _para(doc, [(True, "3. Name of Assessment Center Manager/Officer:  "),
                (False, _dots(45))], size=11)

    # ── Candidate table (bordered) ────────────────────────────────────────────
    headers = ["S/N", "Candidate's\nRegistration Code", "Candidate's Name",
               "Internal Assessor's Marks\n(All Candidates) (100%)",
               "External Verifier's\n(Sampled Candidates) (100%)",
               "Moderated Marks\n(100%)"]
    table = doc.add_table(rows=1 + len(candidates), cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(0.45), Inches(1.55), Inches(1.55),
              Inches(1.30), Inches(1.30), Inches(1.12)]

    hdr = table.rows[0].cells
    for c, label in enumerate(headers):
        _cell(hdr[c], [(False, label)], size=11, bold_all=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, shade=_GREY)
    _row_height(table.rows[0], 42)      # template header row 14

    for i, cand in enumerate(candidates, 1):
        cells = table.rows[i].cells
        _cell(cells[0], [(False, str(cand["sn"]))], align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell(cells[1], [(False, cand["reg_no"])])
        _cell(cells[2], [(False, cand["name"])])
        for c in (3, 4, 5):
            _cell(cells[c], [(False, "")], align=WD_ALIGN_PARAGRAPH.CENTER)
        _row_height(table.rows[i], 18)  # template data rows 15+

    for row in table.rows:
        for c, w in enumerate(widths):
            row.cells[c].width = w

    doc.add_paragraph()  # spacer

    # ── Signatory block ───────────────────────────────────────────────────────
    _para(doc, [(True, "Internal Assessor Reg. Code/National ID. No.: "), (False, _dots(16)),
                (True, "  Signature: "), (False, _dots(16)),
                (True, "  Date: "), (False, _dots(12))], size=11)
    _para(doc, [(True, "External Verifier Reg. Code/National ID. No.: "), (False, _dots(16)),
                (True, "  Signature: "), (False, _dots(16)),
                (True, "  Date: "), (False, _dots(12))], size=11)
    _para(doc, [(True, "Signature: "), (False, _dots(20)),
                (True, "  Date: "), (False, _dots(16)),
                (True, "  Institutional Stamp: "), (False, _dots(20))], size=11)


# ── Public entry point ────────────────────────────────────────────────────────

def build_summative_per_unit_docx(data: dict) -> list[tuple[str, bytes]]:
    """One summative moderated-practical marks sheet (.docx) per unit.
    Returns [(filename, file_bytes), ...]."""
    results: list[tuple[str, bytes]] = []
    for unit in data["units"]:
        doc = Document()
        _setup_page(doc)
        build_summative_doc(doc, data, unit)
        buf = io.BytesIO()
        doc.save(buf)
        fname = _safe_filename(unit["unit_name"], unit.get("unit_code", "")) + ".docx"
        results.append((fname, buf.getvalue()))
    return results
